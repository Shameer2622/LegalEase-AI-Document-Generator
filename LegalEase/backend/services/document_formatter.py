import html
import io
import re
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF

BASE_DIR = Path(__file__).resolve().parents[2]
LOGO_PATH = BASE_DIR / "assets" / "legal_ease_logo.png"


def sanitize_text(text: str) -> str:
    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u00a0": " ",
        "\u2022": "-",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)

    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _term_items(terms: str) -> list[str]:
    return [item.strip() for item in terms.split(";") if item.strip()]


def format_txt(text: str) -> bytes:
    return sanitize_text(text).encode("utf-8")


def format_html_preview(text: str) -> str:
    clean = html.escape(sanitize_text(text))
    paragraphs = re.split(r"\n\s*\n", clean)
    rendered = []

    for paragraph in paragraphs:
        lines = paragraph.splitlines()
        if not lines:
            continue

        first = lines[0].strip()
        is_heading = (
            len(first) <= 100
            and (
                first.isupper()
                or re.match(r"^(?:\d+[\.\)]|[A-Z][\.\)])\s+", first)
                or first.lower() in {
                    "parties",
                    "definitions",
                    "terms",
                    "confidentiality",
                    "termination",
                    "signatures",
                }
            )
        )

        if is_heading:
            rendered.append(f"<h3>{first}</h3>")
            rest = "<br>".join(lines[1:]).strip()
            if rest:
                rendered.append(f"<p>{rest}</p>")
        else:
            rendered.append(f"<p>{'<br>'.join(lines)}</p>")

    return "\n".join(rendered)


def format_docx(text: str, doc_type: str, terms: str = "") -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    if LOGO_PATH.exists():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(str(LOGO_PATH), width=Inches(1.1))

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(15)

    document.add_paragraph()

    for block in re.split(r"\n\s*\n", sanitize_text(text)):
        block = block.strip()
        if not block:
            continue

        lines = block.splitlines()
        first = lines[0].strip()

        if (
            len(first) <= 100
            and (
                first.isupper()
                or re.match(r"^(?:\d+[\.\)]|[A-Z][\.\)])\s+", first)
            )
        ):
            p = document.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(first)
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)

            if len(lines) > 1:
                body = document.add_paragraph("\n".join(lines[1:]))
                body.paragraph_format.space_after = Pt(6)
        else:
            p = document.add_paragraph(block)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15

    items = _term_items(terms)
    if items:
        document.add_paragraph()
        heading = document.add_paragraph()
        r = heading.add_run("KEY TERMS")
        r.bold = True

        table = document.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term"

        for index, item in enumerate(items, start=1):
            cells = table.add_row().cells
            cells[0].text = str(index)
            cells[1].text = item

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LegalEase - AI-generated draft - Review before use")

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class LegalEasePDF(FPDF):
    def __init__(self, doc_type: str):
        super().__init__()
        self.doc_type = doc_type
        self.set_auto_page_break(auto=True, margin=18)

    def header(self):
        if LOGO_PATH.exists():
            try:
                self.image(str(LOGO_PATH), x=95, y=8, w=20)
            except Exception:
                pass

        self.set_y(31)
        self.set_font("Times", "B", 11)
        self.cell(0, 7, self.doc_type.upper(), align="C")
        self.ln(9)

    def footer(self):
        self.set_y(-15)
        self.set_font("Times", "", 8)
        self.cell(0, 8, "LegalEase - AI-generated draft - Review before use", align="C")


def format_pdf(text: str, doc_type: str, terms: str = "") -> bytes:
    pdf = LegalEasePDF(doc_type)
    pdf.add_page()
    pdf.set_font("Times", size=11)

    for block in re.split(r"\n\s*\n", sanitize_text(text)):
        block = block.strip()
        if not block:
            continue

        lines = block.splitlines()
        first = lines[0].strip()
        is_heading = (
            len(first) <= 100
            and (
                first.isupper()
                or re.match(r"^(?:\d+[\.\)]|[A-Z][\.\)])\s+", first)
            )
        )

        if is_heading:
            pdf.set_font("Times", "B", 12)
            pdf.multi_cell(0, 7, first)
            pdf.set_font("Times", "", 11)
            if len(lines) > 1:
                pdf.multi_cell(0, 6, "\n".join(lines[1:]))
        else:
            pdf.multi_cell(0, 6, block)

        pdf.ln(2)

    items = _term_items(terms)
    if items:
        pdf.set_font("Times", "B", 12)
        pdf.cell(0, 8, "KEY TERMS")
        pdf.ln(8)
        pdf.set_font("Times", "", 10)
        for index, item in enumerate(items, start=1):
            pdf.multi_cell(0, 5, f"{index}. {item}")
            pdf.ln(1)

    return bytes(pdf.output())


def export_document(
    file_format: str,
    content: str,
    doc_type: str,
    terms: str = "",
) -> tuple[bytes, str, str]:
    file_format = file_format.lower()

    if file_format == "txt":
        return format_txt(content), "text/plain; charset=utf-8", ".txt"

    if file_format == "docx":
        return (
            format_docx(content, doc_type, terms),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ".docx",
        )

    if file_format == "pdf":
        return format_pdf(content, doc_type, terms), "application/pdf", ".pdf"

    raise ValueError("Unsupported format. Use txt, docx, or pdf.")
